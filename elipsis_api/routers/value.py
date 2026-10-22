"""Client-value routes: assumptions, evidence, action plans, exports, trust."""

from __future__ import annotations

import csv
import io

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, Request, UploadFile
from fastapi.responses import RedirectResponse, Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from constants import (
    PILLAR_LABELS,
)
from elipsis_api import benchmarks, models, services, services_ops, services_webhooks
from elipsis_api.database import get_db
from elipsis_api.deps import get_current_user, login_redirect, templates
from elipsis_api.narrative import board_narrative

router = APIRouter()

_ALLOWED_EVIDENCE = {
    "pdf", "png", "jpg", "jpeg", "webp", "docx", "xlsx", "csv", "txt",
}
_MAX_EVIDENCE = 5 * 1024 * 1024


def _owned(db: Session, user, assessment_id: int):
    assessment = db.get(models.Assessment, assessment_id)
    if assessment is None or not services.can_access_org(db, user, assessment.organization_id):
        return None
    return assessment


def _context(db: Session, assessment: models.Assessment, result: dict, user) -> dict:
    org = db.get(models.Organization, assessment.organization_id)
    dept = db.get(models.Department, assessment.department_id) if assessment.department_id else None
    previous = services.previous_assessment(db, assessment)
    delta = None
    if previous is not None and previous.readiness_index is not None:
        delta = round((result.get("readiness_index") or 0) - previous.readiness_index, 1)
    return {
        "user": user, "assessment": assessment, "result": result, "org": org,
        "organization": org, "department": dept,
        "assumptions": services.get_assumptions(db, assessment.id),
        "benchmark": benchmarks.cohort_benchmark(
            db, org, score=result.get("readiness_index")),
        "reference_points": benchmarks.REFERENCE_POINTS,
        "sources": benchmarks.SOURCES,
        "method_note": benchmarks.METHOD_NOTE,
        "previous": previous, "delta": delta,
    }


@router.get("/assessment/{assessment_id}/assumptions")
def assumptions_get(assessment_id: int, request: Request,
                    user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    _, result = services.compute(db, assessment_id)
    return templates.TemplateResponse(
        request, "assumptions.html", _context(db, assessment, result, user))


@router.post("/assessment/{assessment_id}/assumptions")
def assumptions_post(assessment_id: int, request: Request,
                     staff: int = Form(...),
                     hourly_rate: float = Form(...),
                     working_hours_per_month: float = Form(173.33),
                     max_savings_share: float = Form(12),
                     annual_revenue: str = Form(""),
                     note: str = Form(""),
                     user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    if _owned(db, user, assessment_id) is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    services.save_assumptions(
        db, assessment_id, staff=staff, hourly_rate=hourly_rate,
        working_hours_per_month=working_hours_per_month,
        max_savings_share=max_savings_share / 100.0,
        annual_revenue=annual_revenue, note=note, user_id=user.id)
    assessment = db.get(models.Assessment, assessment_id)
    services_ops.record_audit(
        db, action="assumptions.update",
        organization_id=assessment.organization_id if assessment else None,
        assessment_id=assessment_id, user_id=user.id,
        detail=f"staff={staff}, rate={hourly_rate}, cap={max_savings_share}%")
    db.commit()
    return RedirectResponse(f"/assessment/{assessment_id}/results?assumptions=saved", status_code=303)


@router.post("/assessment/{assessment_id}/evidence/{question_id}")
async def evidence_upload(assessment_id: int, question_id: int, request: Request,
                          file: UploadFile = File(...),
                          user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    if _owned(db, user, assessment_id) is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    question = db.get(models.Question, question_id)
    if question is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    filename = (file.filename or "evidence").replace("\\", "/").split("/")[-1][:180]
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in _ALLOWED_EVIDENCE:
        return RedirectResponse(
            f"/assessment/{assessment_id}/step/1?evidence=type", status_code=303)
    data = await file.read(_MAX_EVIDENCE + 1)
    if len(data) > _MAX_EVIDENCE:
        return RedirectResponse(
            f"/assessment/{assessment_id}/step/1?evidence=size", status_code=303)
    db.add(models.EvidenceAttachment(
        assessment_id=assessment_id, question_id=question_id,
        filename=filename, content_type=file.content_type,
        size_bytes=len(data), data=data, uploaded_by=user.id))
    assessment = db.get(models.Assessment, assessment_id)
    services_ops.record_audit(
        db, action="evidence.upload",
        organization_id=assessment.organization_id if assessment else None,
        assessment_id=assessment_id, user_id=user.id, detail=filename)
    db.commit()
    next_url = request.query_params.get("next") or f"/assessment/{assessment_id}/step/1"
    if not next_url.startswith(f"/assessment/{assessment_id}"):
        next_url = f"/assessment/{assessment_id}/results"
    return RedirectResponse(next_url, status_code=303)


@router.get("/assessment/{assessment_id}/evidence")
def evidence_page(assessment_id: int, request: Request,
                  user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    questions = services.get_questions(db, assessment.questionnaire_id)
    answers = services.answers_map(db, assessment_id)
    attachments = list(db.scalars(select(models.EvidenceAttachment).where(
        models.EvidenceAttachment.assessment_id == assessment_id
    ).order_by(models.EvidenceAttachment.id.desc())))
    by_question = {}
    for row in attachments:
        by_question.setdefault(row.question_id, []).append(row)
    return templates.TemplateResponse(
        request, "evidence.html",
        {"user": user, "assessment": assessment, "questions": questions,
         "answers": answers, "by_question": by_question,
         "organization": db.get(models.Organization, assessment.organization_id)})


@router.get("/evidence/{attachment_id}")
def evidence_download(attachment_id: int, request: Request,
                      user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    row = db.get(models.EvidenceAttachment, attachment_id)
    if row is None or not services.can_access_org(
            db, user, db.get(models.Assessment, row.assessment_id).organization_id):
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    return Response(
        content=row.data, media_type=row.content_type or "application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{row.filename}"'})


@router.post("/assessment/{assessment_id}/evidence/{attachment_id}/delete")
def evidence_delete(assessment_id: int, attachment_id: int, request: Request,
                    user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    if _owned(db, user, assessment_id) is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    row = db.get(models.EvidenceAttachment, attachment_id)
    if row and row.assessment_id == assessment_id:
        assessment = db.get(models.Assessment, assessment_id)
        services_ops.record_audit(
            db, action="evidence.delete",
            organization_id=assessment.organization_id if assessment else None,
            assessment_id=assessment_id, user_id=user.id, detail=row.filename)
        db.delete(row)
        db.commit()
    return RedirectResponse(f"/assessment/{assessment_id}/results?evidence=deleted", status_code=303)


@router.get("/assessment/{assessment_id}/plan")
def plan_get(assessment_id: int, request: Request,
             user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    _, result = services.compute(db, assessment_id)
    if result is None:
        return RedirectResponse(f"/assessment/{assessment_id}/step/1", status_code=303)
    items = services.ensure_plan_items(db, assessment_id, result)
    context = _context(db, assessment, result, user)
    context.update(items=items)
    return templates.TemplateResponse(request, "plan.html", context)


@router.post("/assessment/{assessment_id}/plan/{item_id}")
def plan_update(assessment_id: int, item_id: int, request: Request,
                background_tasks: BackgroundTasks,
                owner: str = Form(""), due_date: str = Form(""),
                status: str = Form("not_started"), notes: str = Form(""),
                user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    if _owned(db, user, assessment_id) is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    services.update_plan_item(
        db, assessment_id, item_id, owner=owner, due_date=due_date,
        status=status, notes=notes)
    assessment = db.get(models.Assessment, assessment_id)
    services_ops.record_audit(
        db, action="plan.update",
        organization_id=assessment.organization_id if assessment else None,
        assessment_id=assessment_id, user_id=user.id, detail=f"item {item_id}: {status}")
    db.commit()
    if assessment is not None:
        for hook in services_webhooks.matching_endpoints(
                db, assessment.organization_id, "action.updated"):
            background_tasks.add_task(
                services_webhooks.deliver_event, hook.id, "action.updated",
                {"event": "action.updated", "item_id": item_id,
                 "status": status, "owner": owner, "assessment_id": assessment_id})
    return RedirectResponse(f"/assessment/{assessment_id}/plan?saved=1", status_code=303)


@router.get("/assessment/{assessment_id}/narrative")
def narrative_page(assessment_id: int, request: Request,
                   user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    _, result = services.compute(db, assessment_id)
    if result is None:
        return RedirectResponse(f"/assessment/{assessment_id}/step/1", status_code=303)
    items = services.ensure_plan_items(db, assessment_id, result)
    context = _context(db, assessment, result, user)
    narrative = board_narrative(
        context["organization"], context["department"], assessment, result,
        benchmark=context["benchmark"], actions=items)
    context.update(narrative=narrative, items=items[:6])
    return templates.TemplateResponse(request, "narrative.html", context)


@router.get("/assessment/{assessment_id}/summary")
def summary(assessment_id: int, request: Request,
            user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    _, result = services.compute(db, assessment_id)
    if result is None:
        return RedirectResponse(f"/assessment/{assessment_id}/step/1", status_code=303)
    items = services.ensure_plan_items(db, assessment_id, result)
    context = _context(db, assessment, result, user)
    context.update(items=items[:6])
    return templates.TemplateResponse(request, "summary.html", context)


@router.get("/assessment/{assessment_id}/export.xlsx")
def export_xlsx(assessment_id: int, request: Request,
                user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return templates.TemplateResponse(request, "404.html", {"user": user}, status_code=404)
    if not services.company_is_pro(db, assessment.organization_id, user):
        return RedirectResponse("/pro?locked=reports", status_code=303)
    _, result = services.compute(db, assessment_id)
    if result is None:
        return RedirectResponse(f"/assessment/{assessment_id}/step/1", status_code=303)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return RedirectResponse(f"/assessment/{assessment_id}/results?export=unavailable",
                                status_code=303)
    org = db.get(models.Organization, assessment.organization_id)
    dept = db.get(models.Department, assessment.department_id) if assessment.department_id else None
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0F172A")
    ws.append(["Elipsis board pack export"])
    ws.append(["Organisation", org.name if org else ""])
    ws.append(["Department", dept.name if dept else "Organisation-wide"])
    ws.append(["Index", result.get("readiness_index")])
    ws.append(["Band", result.get("maturity_band")])
    ws.append(["Confidence", result.get("confidence_index")])
    fin = result.get("financial") or {}
    ws.append(["Annual savings potential", fin.get("annual_savings")])
    ws.append(["Indicative investment", fin.get("investment")])
    ws.append(["ROI %", fin.get("roi")])
    ws.append(["Payback months", fin.get("payback_months")])
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = head
            cell.fill = fill
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    def sheet(name, headers, rows):
        sh = wb.create_sheet(name)
        sh.append(headers)
        for cell in sh[1]:
            cell.font = head
            cell.fill = fill
        for row in rows:
            sh.append(row)
        for idx, width in enumerate([28, 18, 18, 45, 24], start=1):
            sh.column_dimensions[chr(64 + idx)].width = width

    sheet("Pillars", ["Pillar", "Score"], [
        [PILLAR_LABELS.get(code, code), score] for code, score in (result.get("pillars") or {}).items()])
    sheet("Metrics", ["Metric", "Value", "Kind", "Reading"], [
        [m.get("label"), m.get("value"), m.get("kind"), m.get("reading")]
        for rows in (result.get("metrics_by_pillar") or {}).values() for m in rows])
    sheet("Pain points", ["Severity", "Finding", "Impact %", "Detail"], [
        [p.get("severity"), p.get("title"), p.get("impact"), p.get("detail")]
        for p in (result.get("pain_points") or [])])
    sheet("Roadmap", ["Priority", "Action", "Technology", "Horizon", "ROI %"], [
        [r.get("priority"), r.get("title"), r.get("technology"), r.get("horizon"), r.get("roi")]
        for r in (result.get("recommendations") or [])])
    assumptions = (fin.get("assumptions") or {})
    sheet("Assumptions", ["Assumption", "Value"], [[k, v] for k, v in assumptions.items()])
    sheet("Sources", ["Reference", "Value", "Publisher", "URL"], [
        [s.get("label"), s.get("value"), s.get("publisher"), s.get("url")]
        for s in benchmarks.SOURCES])

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="elipsis-{assessment_id}.xlsx"'})


# --- Public trust and distribution pages ---------------------------------
@router.get("/privacy")
def privacy(request: Request, user=Depends(get_current_user)):
    return templates.TemplateResponse(request, "privacy.html", {"user": user})


@router.get("/terms")
def terms(request: Request, user=Depends(get_current_user)):
    return templates.TemplateResponse(request, "terms.html", {"user": user})


@router.get("/security")
def security(request: Request, user=Depends(get_current_user)):
    return templates.TemplateResponse(request, "security.html", {"user": user})


@router.get("/partners")
def partners(request: Request, user=Depends(get_current_user)):
    return templates.TemplateResponse(request, "partners.html", {"user": user})


@router.get("/case-study")
def case_study(request: Request, user=Depends(get_current_user)):
    return templates.TemplateResponse(request, "case_study.html", {"user": user})


@router.get("/assessment/{assessment_id}/import/template.csv")
def import_template(assessment_id: int, user=Depends(get_current_user),
                    db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return Response(status_code=404)
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["code", "score", "evidence"])
    for q in services.get_questions(db, assessment.questionnaire_id):
        writer.writerow([q.code, "", "none"])
    return Response(
        content=out.getvalue(), media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="elipsis-{assessment_id}-answers.csv"'})


def _import_rows(filename: str, raw: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(v or "").strip().lower() for v in rows[0]]
        return [dict(zip(headers, row, strict=False)) for row in rows[1:]]
    text = raw.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


@router.post("/assessment/{assessment_id}/import")
async def import_answers(assessment_id: int, request: Request,
                         file: UploadFile = File(...),
                         user=Depends(get_current_user), db: Session = Depends(get_db)):
    if user is None:
        return login_redirect()
    assessment = _owned(db, user, assessment_id)
    if assessment is None:
        return Response(status_code=404)
    raw = await file.read()
    if not raw or len(raw) > 2 * 1024 * 1024:
        return RedirectResponse(f"/assessment/{assessment_id}/results?import=invalid", status_code=303)
    try:
        rows = _import_rows(file.filename or "", raw)
    except Exception:
        return RedirectResponse(f"/assessment/{assessment_id}/results?import=invalid", status_code=303)
    by_code = {q.code.lower(): q for q in services.get_questions(db, assessment.questionnaire_id)}
    answers = []
    for row in rows:
        code = str(row.get("code") or row.get("question_code") or row.get("question") or "").strip().lower()
        q = by_code.get(code)
        if q is None:
            continue
        try:
            score = int(float(row.get("score") or 0))
        except (TypeError, ValueError):
            continue
        if score < 0 or score > 4:
            continue
        evidence = str(row.get("evidence") or "none").strip().lower()
        if evidence not in {"none", "declared", "verified"}:
            evidence = "none"
        answers.append({"question_id": q.id, "score": score, "evidence": evidence})
    if not answers:
        return RedirectResponse(f"/assessment/{assessment_id}/results?import=empty", status_code=303)
    services.submit(db, assessment_id, answers)
    services_ops.record_audit(
        db, action="answers.import", organization_id=assessment.organization_id,
        assessment_id=assessment_id, user_id=user.id, detail=f"{len(answers)} answers imported")
    db.commit()
    return RedirectResponse(
        f"/assessment/{assessment_id}/results?import={len(answers)}", status_code=303)
